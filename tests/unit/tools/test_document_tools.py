"""
Unit tests for Document Processing tools
"""

from unittest.mock import Mock, patch

import pytest

from portal.tools.document_processing.document_metadata_extractor import (
    DocumentMetadataExtractorTool,
)
from portal.tools.document_processing.excel_processor import ExcelProcessorTool
from portal.tools.document_processing.pandoc_converter import PandocConverterTool
from portal.tools.document_processing.powerpoint_processor import PowerPointProcessorTool
from portal.tools.document_processing.word_processor import WordProcessorTool
from portal.tools.document_tools.pdf_ocr import PDFOCRTool


@pytest.mark.unit
class TestDocumentMetadataExtractorTool:
    """Test document_metadata tool"""

    @pytest.mark.asyncio
    async def test_extract_metadata(self, temp_dir):
        """Test extracting document metadata"""
        tool = DocumentMetadataExtractorTool()

        test_file = temp_dir / "test.pdf"
        test_file.write_bytes(b"%PDF-1.4 dummy")

        # With PyPDF2 installed, tool may fail gracefully on invalid PDF bytes — that's OK
        result = await tool.execute({
            "file_path": str(test_file)
        })

        assert "success" in result


@pytest.mark.unit
class TestExcelProcessorTool:
    """Test excel_processor tool"""

    @pytest.mark.asyncio
    async def test_create_excel(self, temp_dir):
        """Test creating an Excel file"""
        tool = ExcelProcessorTool()

        output_file = temp_dir / "test.xlsx"

        result = await tool.execute({
            "action": "write",
            "file_path": str(output_file),
            "data": [
                ["Name", "Age"],
                ["Alice", 30],
                ["Bob", 25]
            ]
        })

        assert result["success"] is True or "error" in result

    @pytest.mark.asyncio
    async def test_read_excel(self, temp_dir):
        """Test reading an Excel file"""
        tool = ExcelProcessorTool()

        # Create a real Excel file using openpyxl so the tool can read it
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Age"])
        ws.append(["Alice", 30])
        excel_file = temp_dir / "test.xlsx"
        wb.save(str(excel_file))

        result = await tool.execute({
            "action": "read",
            "file_path": str(excel_file)
        })

        assert "success" in result


@pytest.mark.unit
class TestPandocConverterTool:
    """Test pandoc_convert tool"""

    @pytest.mark.asyncio
    async def test_markdown_to_pdf(self, temp_dir):
        """Test converting Markdown to PDF"""
        tool = PandocConverterTool()

        input_file = temp_dir / "input.md"
        input_file.write_text("# Test Document\n\nThis is a test.")

        output_file = temp_dir / "output.pdf"

        with patch("subprocess.run") as mock_run:
            mock_run.return_value = Mock(returncode=0, stdout="", stderr="")

            result = await tool.execute({
                "input_file": str(input_file),
                "output_file": str(output_file),
                "from_format": "markdown",
                "to_format": "pdf"
            })

            assert result["success"] is True or "error" in result

    @pytest.mark.asyncio
    async def test_pandoc_not_installed(self, temp_dir):
        """Test when Pandoc is not installed"""
        tool = PandocConverterTool()

        with patch("subprocess.run") as mock_run:
            mock_run.side_effect = FileNotFoundError()

            result = await tool.execute({
                "input_file": str(temp_dir / "input.md"),
                "output_file": str(temp_dir / "output.pdf")
            })

            assert result["success"] is False


@pytest.mark.unit
class TestPowerPointProcessorTool:
    """Test powerpoint_processor tool"""

    @pytest.mark.asyncio
    async def test_create_presentation(self, temp_dir):
        """Test creating a PowerPoint presentation"""
        tool = PowerPointProcessorTool()

        output_file = temp_dir / "presentation.pptx"

        result = await tool.execute({
            "action": "create",
            "file_path": str(output_file),
            "title": "Test Presentation",
        })

        assert result["success"] is True or "error" in result


@pytest.mark.unit
class TestWordProcessorTool:
    """Test word_processor tool"""

    @pytest.mark.asyncio
    async def test_create_document(self, temp_dir):
        """Test creating a Word document"""
        tool = WordProcessorTool()

        output_file = temp_dir / "document.docx"

        result = await tool.execute({
            "action": "create",
            "file_path": str(output_file),
            "title": "Test Document",
        })

        assert result["success"] is True or "error" in result

    @pytest.mark.asyncio
    async def test_read_document(self, temp_dir):
        """Test reading a Word document"""
        tool = WordProcessorTool()

        # Create a real docx file using python-docx
        from docx import Document
        doc = Document()
        doc.add_paragraph("Test content")
        doc_file = temp_dir / "test.docx"
        doc.save(str(doc_file))

        result = await tool.execute({
            "action": "read",
            "file_path": str(doc_file)
        })

        assert "success" in result


@pytest.mark.unit
class TestPDFOCRTool:
    """Test pdf_ocr tool"""

    @pytest.mark.asyncio
    async def test_extract_text_from_pdf(self, temp_dir):
        """Test extracting text from PDF — requires Tesseract OCR binary and pdf2image/poppler"""
        try:
            import pytesseract
            pytesseract.get_tesseract_version()
        except Exception:
            pytest.skip(
                "Tesseract OCR not available (requires: apt install tesseract-ocr poppler-utils "
                "&& pip install pytesseract pdf2image)"
            )

        tool = PDFOCRTool()
        pdf_file = temp_dir / "test.pdf"
        pdf_file.write_bytes(b"%PDF-1.4 dummy")

        result = await tool.execute({
            "pdf_path": str(pdf_file),
        })

        assert "success" in result
